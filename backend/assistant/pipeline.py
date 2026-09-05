"""
사업자등록증 → 신청서 자동입력 파이프라인 (오케스트레이션).

부품들을 하나의 흐름으로 잇는다:
  이미지/PDF ─▶ [품질 게이트] ─▶ 기본정보 OCR(Qwen) + 업태·종목(EasyOCR)
             ─▶ 구조화 레코드(JSON) ─▶ [신청서 HWPX + 필드매핑] ─▶ fill_hwpx ─▶ 완성 HWPX

공개 함수:
  load_models()                         → Qwen (model, processor) 1회 로딩
  extract_record(src, models, ...)      → 구조화 레코드 dict
  fill_application(record, form, xlsx, out) → 완성 HWPX 경로 + 채워진 필드 로그
  run(src, form, xlsx, out, models)      → 위 둘을 이어서

HWP(구형) 입력은 먼저 HWPX로 변환해야 한다 (backend/assistant/hwpx_convert 예정, Node).
지금은 HWPX 신청서 양식을 직접 받는다.
"""

import io
import json
import os
import tempfile

from backend.assistant.biz_cert_ocr import extract_biz_cert
from backend.assistant import category_ocr
from backend.assistant.biz_cert_quality import check_image_quality

MODEL_ID = "Qwen/Qwen2.5-VL-3B-Instruct"

# 사업자등록증에서 OCR로 뽑는 기본정보 영어 key (필드매핑 xlsx C열과 일치)
_BASIC_KEYS = [
    "trade_name", "corp_name", "ceo_name", "biz_no", "corp_no",
    "birth_date", "open_date", "address_basic",
]

# 법인만: '본점 소재지' 를 따로 읽는 프롬프트 (사업장 소재지와 별개 DB 컬럼)
_HEAD_OFFICE_PROMPT = (
    "이 사업자등록증에는 주소 항목이 두 개 있다: '사업장 소재지' 와 '본점 소재지'. "
    "(내용이 같을 수도, 다를 수도 있다.) 각각의 전체 주소를 정확히 읽어라. "
    "한글로만 적고 한자는 쓰지 마라. 설명하지 말고 JSON만 출력하라.\n"
    '{"사업장소재지": "", "본점소재지": ""}'
)


# ── 모델 로딩 ────────────────────────────────────────────────────────────
def load_models(dtype="bf16"):
    """Qwen2.5-VL 로딩 (bf16 + greedy). EasyOCR reader는 category_ocr에서 lazy 로딩."""
    import torch
    from transformers import AutoProcessor, Qwen2_5_VLForConditionalGeneration

    td = {"bf16": torch.bfloat16, "fp16": torch.float16, "fp32": torch.float32}[dtype]
    model = Qwen2_5_VLForConditionalGeneration.from_pretrained(
        MODEL_ID, torch_dtype=td, device_map="auto"
    )
    processor = AutoProcessor.from_pretrained(MODEL_ID)
    model.eval()
    model.generation_config.do_sample = False
    return model, processor


# ── 입력 정규화 ──────────────────────────────────────────────────────────
def _to_image_path(src, tmp_files):
    """경로/PIL/bytes → (Qwen에 넘길 이미지 경로). PDF는 첫 페이지 렌더링."""
    from PIL import Image, ImageOps

    if hasattr(src, "size"):  # PIL Image
        im = src
    elif isinstance(src, (bytes, bytearray)):
        head = bytes(src[:4])
        if head == b"%PDF":
            im = _pdf_first_page(io.BytesIO(src))
        else:
            im = Image.open(io.BytesIO(src))
    else:  # 경로
        if str(src).lower().endswith(".pdf"):
            im = _pdf_first_page(src)
        else:
            im = Image.open(src)

    im = ImageOps.exif_transpose(im)
    if im.mode != "RGB":
        im = im.convert("RGB")
    fd, path = tempfile.mkstemp(suffix=".png", prefix="bizcert_pipe_")
    os.close(fd)
    im.save(path)
    tmp_files.append(path)
    return path


def _pdf_first_page(path_or_stream, dpi=220):
    import pymupdf
    from PIL import Image

    doc = pymupdf.open(stream=path_or_stream.read(), filetype="pdf") if hasattr(path_or_stream, "read") \
        else pymupdf.open(path_or_stream)
    try:
        pix = doc.load_page(0).get_pixmap(dpi=dpi)
        return Image.open(io.BytesIO(pix.tobytes("png")))
    finally:
        doc.close()


# ── 추출 ────────────────────────────────────────────────────────────────
def extract_record(src, models, *, check_quality=True, refine_category=False, debug=False):
    """이미지/PDF/bytes → 구조화 레코드.

    반환:
      {"ok": bool, "quality": {...},
       "entity_type": "개인"|"법인",
       "trade_name":..., ..., "address_basic":...,
       "business_category": [{"업태": str, "종목": [str, ...]}],
       "meta": {"category_needs_review": bool, "category_review_reason": str}}
    실패(품질 미달) 시: {"ok": False, "quality": {...}}
    """
    model, processor = models
    tmp_files = []
    try:
        img_path = _to_image_path(src, tmp_files)

        quality = check_image_quality(img_path) if check_quality else {"acceptable": True, "score": None}
        if check_quality and not quality["acceptable"]:
            return {"ok": False, "quality": quality}

        # 1) 기본정보 (Qwen)
        entity_type, biz_cert = extract_biz_cert(img_path, model, processor)

        # 2) 업태·종목 (EasyOCR + 좌표규칙, 옵션으로 Qwen 텍스트 보정)
        qwen = (model, processor) if refine_category else None
        groups = category_ocr.extract_categories(img_path, qwen=qwen, debug=debug)
        needs_review, reason = category_ocr.is_low_confidence(groups)

        # 3) 법인이면 '본점 소재지' 도 별도로 (사업장 소재지 = address_basic 와 다른 DB 컬럼)
        head_office = ""
        if entity_type == "법인":
            try:
                from backend.assistant.biz_cert_ocr import ask_image, parse_json
                from backend.assistant.biz_cert_ocr import load_image as _li
                raw = ask_image(model, processor, _li(img_path), _HEAD_OFFICE_PROMPT)
                p = parse_json(raw) or {}
                head_office = str(p.get("본점소재지") or "").strip()
                # 본점을 못 읽었으면 이 호출의 사업장소재지로도 보정 시도
                if not head_office:
                    head_office = str(p.get("사업장소재지") or "").strip()
            except Exception:
                head_office = ""

        record = {
            "ok": True,
            "quality": quality,
            "entity_type": entity_type,
            "business_category": groups,
            "head_office_address": head_office,  # 본점 소재지 (법인만, 없으면 "")
            "meta": {"category_needs_review": needs_review, "category_review_reason": reason},
        }
        record.update({k: biz_cert.get(k, "") for k in _BASIC_KEYS})
        return record
    finally:
        for f in tmp_files:
            try:
                os.remove(f)
            except OSError:
                pass


# ── 신청서 채우기 ────────────────────────────────────────────────────────
def _load_mapping(xlsx_path):
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path)
    ws = next((wb[n] for n in wb.sheetnames if "매핑" in n), wb.active)
    rules = []
    for r in range(2, ws.max_row + 1):
        key, kw = ws.cell(r, 3).value, ws.cell(r, 6).value
        if key and kw:
            rules.append({"key": str(key).strip(),
                          "keywords": [k.strip() for k in str(kw).split(",") if k.strip()]})
    return rules


def _record_to_biz_cert(record):
    """레코드 → fill_hwpx가 쓰는 (biz_cert dict, entity_type).

    업태·종목은 다중값이라, 대표(첫 번째) 업태-종목을 단일 문자열로 매핑한다.
    필드매핑 xlsx의 key: biz_type=업태, biz_item=종목.
    """
    bc = {k: record.get(k, "") for k in _BASIC_KEYS}
    entity = record.get("entity_type", "개인")
    # 신청서 '본사 주소' 칸: 법인=본점 소재지, 개인=사업장 소재지
    #   공장 주소는 사업자등록증으로 판정 불가 → hwpx_fill 이 '공장' 라벨을 항상 비움
    if entity == "법인" and record.get("head_office_address"):
        bc["address_basic"] = record["head_office_address"]
    cats = record.get("business_category") or []
    if cats:
        bc["biz_type"] = cats[0].get("업태", "")
        bc["biz_item"] = ", ".join(cats[0].get("종목", []))
    return bc, entity


def fill_application(record, form_hwpx, mapping_xlsx, out_path, models=None):
    """구조화 레코드 + 신청서 HWPX + 필드매핑 → 완성 HWPX. (채워진 필드 로그 반환)

    fill_hwpx_all: 모든 섹션 순회 + 중복 라벨 다 채움 + 서식 태그 제거 매칭
    (팀원 원본 biz_cert_ocr.fill_hwpx 의 보강판).
    models 를 넘기면 '대표자' 계열 라벨 판별에 LangChain(로컬 Qwen) 체인을 쓴다.
    """
    from backend.assistant.hwpx_fill import fill_hwpx_all

    rules = _load_mapping(mapping_xlsx)
    biz_cert, entity_type = _record_to_biz_cert(record)
    log = fill_hwpx_all(form_hwpx, out_path, rules, biz_cert, entity_type, models=models)
    return {"out": out_path, "filled": log}


# ── 전체 ────────────────────────────────────────────────────────────────
def run(src, form_hwpx, mapping_xlsx, out_path, models, **kw):
    """사업자등록증 이미지 → 완성된 신청서 HWPX 까지 한 번에."""
    record = extract_record(src, models, **kw)
    if not record["ok"]:
        return {"ok": False, "quality": record["quality"]}
    fill = fill_application(record, form_hwpx, mapping_xlsx, out_path, models=models)
    return {"ok": True, "record": record, **fill}


if __name__ == "__main__":
    import sys

    args = sys.argv[1:]
    if len(args) < 3:
        print("usage: python -m backend.assistant.pipeline <biz_cert_img> <form.hwpx> <mapping.xlsx> [out.hwpx]")
        sys.exit(1)
    src, form, xlsx = args[:3]
    out = args[3] if len(args) > 3 else "완성_신청서.hwpx"

    print("모델 로딩...")
    m = load_models()
    print("완료\n")
    res = run(src, form, xlsx, out, m, debug=True)
    if not res["ok"]:
        print("품질 미달:", res["quality"])
        sys.exit(2)
    print("\n=== 레코드 ===")
    print(json.dumps(res["record"], ensure_ascii=False, indent=2))
    print("\n=== 채워진 필드 ===")
    for f, v in res["filled"]:
        print(f"  [{f}] <- {v}")
    print(f"\n완성 -> {res['out']}")
